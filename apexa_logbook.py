"""The experiment logbook — read it where it lives, including over SSH.

During a beamtime the team keeps a running log: what was changed, what broke,
what the 3am decision was. APEXA could not see any of it, so neither could a
report.

Two backends behind one interface:

    file   a path, local or on another host, read over SSH     works air-gapped
    web    a shared document over HTTP                         needs the web tier

The file backend is the default because it is the only one that works on the
primary deployment: APEXA runs on copland, which has the ANL internal network
and no public internet, and SSHes to whichever machine holds the log.

**Why this does not reuse remote_read_text.** That helper is `head -c` -- the
FIRST bytes -- which is correct for a parameter file and precisely wrong here.
A logbook grows all day; its first 256 KiB is the stalest part of it. Reading a
long log from the head would show yesterday's setup and silently omit the
entries someone is asking about. So the file backend tails.

Nothing here summarises. It returns entries verbatim with their timestamps;
whether a summary is written alongside them is the report's decision, and the
report labels it as generated. A paraphrase that replaces its source is the
failure mode the ledger exists to prevent.
"""
from __future__ import annotations

import os
import re
import shlex
import tempfile
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

__all__ = [
    "LogbookEntry", "LogbookRead", "LogbookSource",
    "FileLogbookSource", "WebLogbookSource",
    "resolve_logbook_source", "parse_entries",
]

DEFAULT_TAIL_BYTES = 262_144
DEFAULT_ENTRIES = 12
_OFFICE = (".docx", ".xlsx", ".odt", ".ods", ".pdf")


# ── entries ───────────────────────────────────────────────────────────────────
@dataclass
class LogbookEntry:
    text: str
    timestamp: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class LogbookRead:
    entries: List[LogbookEntry] = field(default_factory=list)
    source: str = ""
    host: Optional[str] = None
    truncated: bool = False          # the log was longer than we tailed
    bytes_read: int = 0
    error: Optional[str] = None

    @property
    def ok(self) -> bool:
        return self.error is None

    def to_dict(self) -> Dict[str, Any]:
        d = asdict(self)
        d["entries"] = [e.to_dict() for e in self.entries]
        d["ok"] = self.ok
        return d


# A line that opens a new entry. Logbooks are written by people, so accept the
# shapes people actually use rather than imposing one.
_TS_PATTERNS = (
    r"\d{4}-\d{2}-\d{2}[ T]\d{1,2}:\d{2}(?::\d{2})?",   # 2026-09-26 14:32[:05]
    r"\d{4}/\d{2}/\d{2}\s+\d{1,2}:\d{2}",               # 2026/09/26 14:32
    r"\d{1,2}:\d{2}(?::\d{2})?\s*(?:[AaPp][Mm])?",       # 14:32 / 2:32 pm
    r"\d{4}-\d{2}-\d{2}",                               # 2026-09-26
)
_ENTRY_RE = re.compile(
    r"^\s*(?:[-*>]\s*)?(?:[\[(]\s*)?(" + "|".join(_TS_PATTERNS) + r")\s*(?:[\])]|[-–—:])?\s*",
)
_HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+(.*\S)\s*$")


def parse_entries(text: str) -> List[LogbookEntry]:
    """Split a logbook into entries, newest last.

    Tries, in order: timestamped lines, markdown headings, blank-line blocks.
    Falls back to the whole text as one entry -- an unparseable log is still
    worth showing verbatim, and inventing structure would misattribute
    timestamps to the wrong observation.
    """
    if not text or not text.strip():
        return []
    lines = text.splitlines()

    starts: List[Tuple[int, Optional[str]]] = []
    for i, ln in enumerate(lines):
        m = _ENTRY_RE.match(ln)
        if m and ln.strip():
            starts.append((i, m.group(1)))
    if not starts:
        for i, ln in enumerate(lines):
            h = _HEADING_RE.match(ln)
            if h:
                starts.append((i, None))

    if starts:
        out: List[LogbookEntry] = []
        for n, (i, ts) in enumerate(starts):
            j = starts[n + 1][0] if n + 1 < len(starts) else len(lines)
            body = "\n".join(lines[i:j]).strip()
            if body:
                out.append(LogbookEntry(text=body, timestamp=ts))
        return out

    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    if len(blocks) > 1:
        return [LogbookEntry(text=b) for b in blocks]
    return [LogbookEntry(text=text.strip())]


# ── sources ───────────────────────────────────────────────────────────────────
class LogbookSource:
    """Where the logbook is read from."""

    kind = "abstract"

    def read(self, n_entries: int = DEFAULT_ENTRIES) -> LogbookRead:  # pragma: no cover
        raise NotImplementedError


class FileLogbookSource(LogbookSource):
    """A file, local or on another host over SSH. The default."""

    kind = "file"

    def __init__(self, path: str, host: str = "", tail_bytes: int = DEFAULT_TAIL_BYTES):
        self.path = str(path)
        self.host = (host or "").strip()
        self.tail_bytes = int(tail_bytes)

    # -- local ----------------------------------------------------------------
    def _read_local(self) -> Tuple[str, bool, int]:
        p = Path(self.path).expanduser()
        if not p.is_file():
            raise FileNotFoundError(f"no logbook at {p}")
        if p.suffix.lower() in _OFFICE:
            return _office_text(p), False, p.stat().st_size
        size = p.stat().st_size
        with open(p, "rb") as fh:
            if size > self.tail_bytes:
                fh.seek(-self.tail_bytes, os.SEEK_END)
            raw = fh.read()
        return raw.decode("utf-8", errors="replace"), size > self.tail_bytes, len(raw)

    # -- remote ---------------------------------------------------------------
    def _read_remote(self) -> Tuple[str, bool, int]:
        from apexa_remote_exec import remote_run, remote_get, ssh_hint

        q = shlex.quote(self.path)
        if Path(self.path).suffix.lower() in _OFFICE:
            # Binary: copy it over, then parse locally. tail -c on a .docx is
            # a truncated zip and parses as nothing.
            tmp = Path(tempfile.mkdtemp()) / Path(self.path).name
            g = remote_get(self.host, self.path, str(tmp))
            if not g.get("success"):
                raise RuntimeError(g.get("stderr") or ssh_hint(self.host))
            return _office_text(tmp), False, tmp.stat().st_size

        # Size first, so "truncated" is a fact rather than a guess at whether
        # the tail happened to fill.
        s = remote_run(self.host, f"wc -c < {q}", timeout=20)
        size = 0
        if s.get("success"):
            try:
                size = int((s.get("stdout") or "0").strip() or 0)
            except ValueError:
                size = 0
        r = remote_run(self.host, f"tail -c {self.tail_bytes} {q}", timeout=60)
        if not r.get("success"):
            raise RuntimeError(
                ssh_hint(self.host) if r.get("ssh_failed")
                else (r.get("stderr") or f"could not read {self.path} on {self.host}"))
        text = r.get("stdout", "")
        return text, bool(size and size > self.tail_bytes), len(text.encode(errors="replace"))

    def read(self, n_entries: int = DEFAULT_ENTRIES) -> LogbookRead:
        out = LogbookRead(source=self.path, host=self.host or None)
        try:
            text, truncated, nbytes = (self._read_remote() if self.host
                                       else self._read_local())
        except Exception as e:
            out.error = f"{type(e).__name__}: {e}"
            return out
        out.truncated, out.bytes_read = truncated, nbytes
        entries = parse_entries(text)
        if truncated and entries:
            # The tail almost certainly began mid-entry; that fragment would be
            # shown with the wrong timestamp attached.
            entries = entries[1:]
        out.entries = entries[-n_entries:] if n_entries > 0 else entries
        return out


class WebLogbookSource(LogbookSource):
    """A shared document over HTTP. Additive: only on a web-tier host.

    Deliberately not wired to a specific provider. A real Google Doc needs
    OAuth, and storing those secrets on a shared, unauthenticated beamline host
    is a decision for the operator, not a default. Point this at any URL that
    returns the log as text (a published doc, an export endpoint, a wiki raw
    view) and it works with no credentials at all.
    """

    kind = "web"

    def __init__(self, url: str, timeout: int = 30, headers: Optional[Dict[str, str]] = None):
        self.url, self.timeout, self.headers = url, timeout, dict(headers or {})

    def read(self, n_entries: int = DEFAULT_ENTRIES) -> LogbookRead:
        out = LogbookRead(source=self.url)
        try:
            import apexa_network as net
        except ImportError:
            net = None                # standalone use of this module: no gate to apply
        if net is not None:
            # Fail CLOSED. An earlier version wrapped this in `except Exception:
            # pass` and called a function that had been renamed -- the gate
            # silently vanished and the request went out, which on an
            # air-gapped host is the blocking hang this tier system exists to
            # prevent. If the module is present but cannot answer, do not guess.
            try:
                allowed = net.has_web()
                where = net.tier()
            except Exception as e:
                out.error = (f"cannot determine the network tier ({e}); refusing to "
                             f"reach the network. Use a file logbook here.")
                return out
            if not allowed:
                out.error = (
                    f"a web logbook needs APEXA_NETWORK=web; this host is at "
                    f"'{where}'. Use a file logbook (local or over SSH) here, or "
                    f"export the document to a machine APEXA can reach.")
                return out
        try:
            import urllib.request
            req = urllib.request.Request(self.url, headers=self.headers)
            with urllib.request.urlopen(req, timeout=self.timeout) as r:
                raw = r.read()
        except Exception as e:
            out.error = f"{type(e).__name__}: {e}"
            return out
        text = raw.decode("utf-8", errors="replace")
        out.bytes_read = len(raw)
        entries = parse_entries(text)
        out.entries = entries[-n_entries:] if n_entries > 0 else entries
        return out


def _office_text(path: Path) -> str:
    """Text out of a .docx/.xlsx/.pdf logbook. Lazy imports, clear failures."""
    ext = path.suffix.lower()
    if ext == ".docx":
        try:
            import docx  # type: ignore
        except ImportError:
            raise RuntimeError("reading a .docx logbook needs python-docx")
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append("  ".join(cells))
        return "\n".join(parts)
    if ext in (".xlsx", ".ods"):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            raise RuntimeError("reading a spreadsheet logbook needs openpyxl")
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    rows.append("  ".join(cells))
        return "\n".join(rows)
    if ext == ".pdf":
        try:
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader  # type: ignore
        except ImportError:
            raise RuntimeError("reading a .pdf logbook needs pypdf")
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    return path.read_text(errors="replace")


# ── resolution ────────────────────────────────────────────────────────────────
def resolve_logbook_source(spec: str = "", *, host: str = "",
                           tail_bytes: int = DEFAULT_TAIL_BYTES
                           ) -> Optional[LogbookSource]:
    """Turn a spec into a source.

        /data/logbook.md              local file
        copland:/data/logbook.md      that file, over SSH
        https://…/export?format=txt   web backend
        (empty)                       $APEXA_LOGBOOK, else None

    A Windows drive letter is not a host, so `C:\\log.md` stays local.
    """
    spec = (spec or os.environ.get("APEXA_LOGBOOK", "")).strip()
    if not spec:
        return None
    if spec.startswith(("http://", "https://")):
        return WebLogbookSource(spec)
    h = (host or os.environ.get("APEXA_LOGBOOK_HOST", "")).strip()
    if not h and ":" in spec:
        head, _, tail = spec.partition(":")
        if head and tail and not (len(head) == 1 and head.isalpha()):
            h, spec = head, tail
    return FileLogbookSource(spec, host=h, tail_bytes=tail_bytes)
