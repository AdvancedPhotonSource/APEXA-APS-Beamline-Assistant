"""Place the new slides in the narrative and correct counts the deck has outgrown."""
from pptx import Presentation

prs = Presentation("2610-Genesis-Demo.NEW.pptx")

# ── 1. correct stale counts ────────────────────────────────────────────────
# Measured now: core 11, midas 55, motor 13, gsas2 5 = 84 across 4 servers.
FIX = {
    "56 Tools Across 3 MCP Servers (Current Deployment)":
        "84 Tools Across 4 MCP Servers (Current Deployment)",
    "Core Server (9 tools)":   "Core Server (11 tools)",
    "MIDAS Server (34 tools)": "MIDAS Server (55 tools)",
    "Motor Server (13 tools)": "Motor Server (13 tools)",
}
changed = []
for i, s in enumerate(prs.slides):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                t = run.text.strip()
                if t in FIX and FIX[t] != t:
                    run.text = run.text.replace(t, FIX[t])
                    changed.append((i, t, FIX[t]))
for c in changed:
    print(f"  slide {c[0]}: {c[1]!r} -> {c[2]!r}")

# ── 2. move the five new slides into the narrative ─────────────────────────
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)

def title_of(idx):
    s = prs.slides[idx]
    return s.shapes.title.text if s.shapes.title is not None else ""

new = ids[37:42]                      # the five appended slides, in order
for e in new:
    sldIdLst.remove(e)

remaining = list(sldIdLst)
def index_of_title(prefix):
    for n, e in enumerate(remaining):
        # resolve element -> slide via position in the (now shortened) list
        s = prs.slides[n]
        t = s.shapes.title.text if s.shapes.title is not None else ""
        if t.startswith(prefix):
            return n
    return None

# "the cost" + "FF-HEDM order" go right after the problem statement
anchor_problem = 3
for off, e in enumerate(new[:2]):
    sldIdLst.insert(anchor_problem + 1 + off, e)

# the three capability/example slides follow the existing calibration use case
remaining = list(sldIdLst)
anchor_use = None
for n in range(len(remaining)):
    t = prs.slides[n].shapes.title.text if prs.slides[n].shapes.title is not None else ""
    if t.startswith("Example Use Case"):
        anchor_use = n
        break
if anchor_use is None:
    anchor_use = len(remaining) - 1
for off, e in enumerate(new[2:]):
    sldIdLst.insert(anchor_use + 1 + off, e)

prs.save("2610-Genesis-Demo.NEW.pptx")

prs2 = Presentation("2610-Genesis-Demo.NEW.pptx")
print(f"\nfinal order ({len(prs2.slides)} slides):")
for i, s in enumerate(prs2.slides):
    t = s.shapes.title.text if s.shapes.title is not None else ""
    if not t:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().split("\n")[0]; break
    mark = "  <== NEW" if t.startswith(("What a MIDAS", "Example: FF-HEDM",
                                        "Example: Forward Simulation",
                                        "DFXM in APEXA", "Why the Guards")) else ""
    print(f"{i:2d}  {t[:66]}{mark}")
